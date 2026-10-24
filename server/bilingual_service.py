import os
from pathlib import Path
from typing import Dict, Any, Optional

import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from server.ai_client import ai_client
from server.export.xml_helpers import (
    set_page_borders,
    set_cell_borders,
    set_cell_shading,
    set_section_columns,
    add_footer_with_page_number,
    configure_document_fonts,
    set_paragraph_bidi,
    set_run_bidi
)

EXPORT_DIR = Path(__file__).resolve().parent.parent / "data"
LOGO_PATH = Path(__file__).resolve().parent.parent / "assets" / "logo.jpg"
LOGO_WATERMARK_PATH = Path(__file__).resolve().parent.parent / "assets" / "logo_watermark.png"

def translate_paper_to_bilingual(paper_data: Dict[str, Any]) -> Dict[str, Any]:
    """
    Translates an existing paper into a bilingual format (English ⇄ Urdu).
    Uses ai_client (Cloudflare Worker online or local offline engine).
    """
    bilingual_paper = json_clone(paper_data)
    bilingual_paper["is_bilingual"] = True

    for group in bilingual_paper.get("groups", []):
        # Translate group heading
        raw_heading = group.get("title", "")
        trans_res = ai_client.call(
            task="translate",
            payload={"text": raw_heading, "source_lang": "en", "target_lang": "ur"}
        )
        group["urdu_title"] = trans_res.get("translated_text", raw_heading)

        # Translate questions
        for q in group.get("questions", []):
            q_text = q.get("text", "")
            q_trans = ai_client.call(
                task="translate",
                payload={"text": q_text, "source_lang": "en", "target_lang": "ur"}
            )
            q["urdu_text"] = q_trans.get("translated_text", q_text)

            # Translate options if MCQ
            if q.get("options"):
                urdu_opts = []
                for opt in q["options"]:
                    opt_trans = ai_client.call(
                        task="translate",
                        payload={"text": opt, "source_lang": "en", "target_lang": "ur"}
                    )
                    urdu_opts.append(opt_trans.get("translated_text", opt))
                q["urdu_options"] = urdu_opts

    return bilingual_paper

def build_bilingual_docx(paper: Dict[str, Any], output_path: Optional[Path] = None) -> Path:
    """
    Generates a Word 2007-compatible bilingual question paper (.docx)
    combining English (Times New Roman) and Urdu (Jameel Noori Nastaleeq RTL).
    """
    header = paper.get("header", {})
    template = paper.get("template", "3a")
    class_name = header.get("class_name", "Class 9th")
    subject = header.get("subject", "Social Science")
    paper_name = header.get("paper_name", "FIRST UNIT TEST").upper()
    month_year = header.get("month_year", "").upper()
    exam_title = f"{paper_name} — {month_year}" if month_year else paper_name

    clean_class = class_name.replace("Class", "").strip()
    filename = f"Class {clean_class} {subject} - Bilingual Paper.docx"
    final_output = output_path or (EXPORT_DIR / filename)

    doc = Document()
    configure_document_fonts(doc)

    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    # 12mm margins
    section.top_margin = Inches(0.47)
    section.bottom_margin = Inches(0.47)
    section.left_margin = Inches(0.47)
    section.right_margin = Inches(0.47)

    set_page_borders(section)

    # Watermark in Header
    if LOGO_WATERMARK_PATH.exists():
        header_para = section.header.paragraphs[0]
        header_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = header_para.add_run()
        run.add_picture(str(LOGO_WATERMARK_PATH), width=Inches(4.5))

    # Page Footer
    footer_para = section.footer.paragraphs[0]
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_para.style.font.size = Pt(8.5)
    footer_para.style.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    prefix = f"{header.get('school', 'Al-Rizwan English School')} · {class_name} · {subject} · {exam_title}"
    add_footer_with_page_number(footer_para, prefix)

    # Masthead Block
    table_masthead = doc.add_table(rows=1, cols=3)
    table_masthead.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_masthead.autofit = False

    # Left: Logo
    cell_logo = table_masthead.rows[0].cells[0]
    cell_logo.width = Inches(1.1)
    set_cell_borders(cell_logo, top=None, bottom="single", left=None, right=None, color="1A1A1A", sz="18")
    p_logo = cell_logo.paragraphs[0]
    p_logo.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if LOGO_PATH.exists():
        p_logo.add_run().add_picture(str(LOGO_PATH), width=Inches(0.95))

    # Center: Titles
    cell_center = table_masthead.rows[0].cells[1]
    cell_center.width = Inches(4.7)
    set_cell_borders(cell_center, top=None, bottom="single", left=None, right=None, color="1A1A1A", sz="18")
    p_title = cell_center.paragraphs[0]
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(2)

    r_sch = p_title.add_run(header.get("school", "Al-Rizwan English School") + "\n")
    r_sch.font.name = "Times New Roman"
    r_sch.font.size = Pt(20)
    r_sch.font.bold = True

    r_addr = p_title.add_run(header.get("address", "Chaitaniya Nagar Road, Nanded") + "\n")
    r_addr.font.name = "Times New Roman"
    r_addr.font.size = Pt(10)

    r_exam = p_title.add_run(exam_title)
    r_exam.font.name = "Times New Roman"
    r_exam.font.size = Pt(10.5)
    r_exam.font.bold = True

    # Right: Roll No Boxes
    cell_roll = table_masthead.rows[0].cells[2]
    cell_roll.width = Inches(1.5)
    set_cell_borders(cell_roll, top=None, bottom="single", left=None, right=None, color="1A1A1A", sz="18")
    p_roll = cell_roll.paragraphs[0]
    p_roll.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_roll_lbl = p_roll.add_run("Roll No.\n")
    r_roll_lbl.font.size = Pt(9)

    tbl_boxes = cell_roll.add_table(rows=1, cols=4)
    tbl_boxes.alignment = WD_TABLE_ALIGNMENT.RIGHT
    for box_cell in tbl_boxes.rows[0].cells:
        box_cell.width = Inches(0.24)
        set_cell_borders(box_cell, "single", "single", "single", "single", "1A1A1A", "8")
        bp = box_cell.paragraphs[0]
        bp.paragraph_format.space_before = Pt(0)
        bp.paragraph_format.space_after = Pt(0)
        brun = bp.add_run(" ")
        brun.font.size = Pt(10)

    # Details Table
    details_table = doc.add_table(rows=3, cols=6)
    details_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    col_w = Inches(7.33 / 6)
    for row in details_table.rows:
        for cell in row.cells:
            cell.width = col_w
            set_cell_borders(cell, "single", "single", "single", "single", "1A1A1A", "8")

    # Row 0
    _set_kv(details_table.cell(0, 0), details_table.cell(0, 1), "Class :", class_name)
    _set_kv(details_table.cell(0, 2), details_table.cell(0, 3), "Subject :", subject)
    _set_kv(details_table.cell(0, 4), details_table.cell(0, 5), "Date :", header.get("date", ""))

    # Row 1
    _set_kv(details_table.cell(1, 0), details_table.cell(1, 1), "Marks :", f"{header.get('marks', 40)}")
    _set_kv(details_table.cell(1, 2), details_table.cell(1, 3), "Exam :", header.get("paper_name", ""))
    _set_kv(details_table.cell(1, 4), details_table.cell(1, 5), "Time :", header.get("time", ""))

    # Row 2: Name
    c_name_start = details_table.cell(2, 0)
    c_name_end = details_table.cell(2, 5)
    c_name_start.merge(c_name_end)
    pn = c_name_start.paragraphs[0]
    pn.paragraph_format.space_after = Pt(2)
    rn_lbl = pn.add_run("Name of the Student : ")
    rn_lbl.font.size = Pt(9.5)
    rn_dot = pn.add_run("." * 65)
    rn_dot.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Render Bilingual Groups & Questions
    for group in paper.get("groups", []):
        # English Group Heading
        p_grp = doc.add_paragraph()
        p_grp.paragraph_format.space_before = Pt(8)
        p_grp.paragraph_format.space_after = Pt(2)
        r_grp = p_grp.add_run(group.get("title", ""))
        r_grp.font.name = "Times New Roman"
        r_grp.font.size = Pt(13)
        r_grp.font.bold = True

        # Urdu Group Heading (RTL)
        if group.get("urdu_title"):
            p_u_grp = doc.add_paragraph()
            set_paragraph_bidi(p_u_grp, align="right")
            p_u_grp.paragraph_format.space_before = Pt(0)
            p_u_grp.paragraph_format.space_after = Pt(4)
            r_u_grp = p_u_grp.add_run(group["urdu_title"])
            set_run_bidi(r_u_grp, "Jameel Noori Nastaleeq")
            r_u_grp.font.size = Pt(13)
            r_u_grp.font.bold = True

        # Questions
        for q in group.get("questions", []):
            q_num = q.get("q_num", "")
            q_text = q.get("text", "")
            marks_str = f" [{int(q['marks'])}]" if q.get("marks") and q["marks"] > 1 else ""

            # English line
            p_q = doc.add_paragraph()
            p_q.paragraph_format.space_before = Pt(2)
            p_q.paragraph_format.space_after = Pt(1)
            rq_num = p_q.add_run(f"{q_num}. ")
            rq_num.font.bold = True
            rq_num.font.size = Pt(11)
            rq_text = p_q.add_run(f"{q_text}{marks_str}")
            rq_text.font.size = Pt(11)

            # Urdu line
            if q.get("urdu_text"):
                p_uq = doc.add_paragraph()
                set_paragraph_bidi(p_uq, align="right")
                p_uq.paragraph_format.space_before = Pt(0)
                p_uq.paragraph_format.space_after = Pt(3)
                ruq = p_uq.add_run(q["urdu_text"])
                set_run_bidi(ruq, "Jameel Noori Nastaleeq")
                ruq.font.size = Pt(12)

            # Bilingual Options for MCQs
            if q.get("options") and len(q["options"]) == 4:
                urdu_opts = q.get("urdu_options", ["", "", "", ""])
                opt_table = doc.add_table(rows=2, cols=2)
                opt_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                w_half = Inches(3.6)

                opt_cells = [
                    opt_table.cell(0, 0), opt_table.cell(0, 1),
                    opt_table.cell(1, 0), opt_table.cell(1, 1)
                ]

                for idx in range(4):
                    cell = opt_cells[idx]
                    cell.width = w_half
                    set_cell_borders(cell, None, None, None, None)
                    p_opt = cell.paragraphs[0]
                    p_opt.paragraph_format.space_after = Pt(1)

                    # English option
                    ro_en = p_opt.add_run(f"({idx+1}) {q['options'][idx]}  /  ")
                    ro_en.font.size = Pt(10)

                    # Urdu option
                    if idx < len(urdu_opts) and urdu_opts[idx]:
                        ro_ur = p_opt.add_run(urdu_opts[idx])
                        set_run_bidi(ro_ur, "Jameel Noori Nastaleeq")
                        ro_ur.font.size = Pt(10.5)

    doc.save(str(final_output))
    return final_output

def _set_kv(c_lbl, c_val, label, value):
    p_lbl = c_lbl.paragraphs[0]
    p_lbl.paragraph_format.space_after = Pt(0)
    r1 = p_lbl.add_run(label)
    r1.font.size = Pt(9.5)

    p_val = c_val.paragraphs[0]
    p_val.paragraph_format.space_after = Pt(0)
    r2 = p_val.add_run(str(value))
    r2.font.bold = True
    r2.font.size = Pt(9.5)

def json_clone(obj):
    import json
    return json.loads(json.dumps(obj))
